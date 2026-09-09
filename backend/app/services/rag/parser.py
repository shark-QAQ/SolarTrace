# app/services/rag/parser.py
"""文档解析：PDF / DOCX / Markdown / TXT → 保留标题结构的段落组。

解析输出与文件格式解耦，分块器消费统一的 Section 结构；
解析失败抛 ParseFailed，由 API 层转为结构化错误。
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field

from app.core.errors import ParseFailed

# 常见条款/章节标题启发式（针对无结构化样式的纯文本/PDF）
_NUM_HEAD = re.compile(r"^\s*(第[一二三四五六七八九十百零\d]+[章节部分篇]|附录\s*[A-Z]?|[一二三四五六七八九十]+[、\.])")
_ARTICLE_HEAD = re.compile(r"^\s*第[一二三四五六七八九十百零\d]+条\b")
_DOC_HEAD = re.compile(r"^\s*\d+(\.\d+)*[、\.\s]")


@dataclass
class Section:
    """一个带标题链的文本区块。path 含所有层级标题。"""

    path: list[str]           # 标题层级链，如 ["第3章 电气", "3.2 直流系统"]
    paragraphs: list[str] = field(default_factory=list)

    @property
    def full_text(self) -> str:
        return "\n".join(p for p in self.paragraphs if p.strip()).strip()


def _append(sections: list[Section], path: list[str]) -> Section:
    for s in sections:
        if s.path == path:
            return s
    s = Section(path=path)
    sections.append(s)
    return s


def parse_text(text: str, default_heading: str = "") -> list[Section]:
    """纯文本/Markdown 解析。

    Markdown ``#`` 层级优先；无结构时按数字/条款启发式断标题。
    """
    sections: list[Section] = []
    current: Section | None = None
    heading_stack: list[str] = []

    def make_section(path: list[str]) -> Section:
        return _append(sections, list(path))

    def flush(path: list[str]) -> Section:
        nonlocal current, heading_stack
        heading_stack = list(path)
        current = make_section(path)
        return current

    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        md = re.match(r"^#{1,6}\s+(.+)$", line.strip())
        if md:
            level = len(md.group(0)) - len(md.group(0).lstrip("#"))
            title = md.group(1).strip()
            heading_stack = heading_stack[: level - 1] + [title]
            current = flush(heading_stack)
            continue
        if current is None:
            current = flush([default_heading] if default_heading else [])
        stripped = line.strip()
        # 文本形态的章节/条款标题启发式
        if (len(stripped) <= 40 and _NUM_HEAD.match(stripped) and not stripped.endswith(("。", "；", "：", ";"))) or (
            _ARTICLE_HEAD.match(stripped)
        ):
            heading_stack = heading_stack[:1] + [stripped] if heading_stack else [stripped]
            current = flush(heading_stack)
            continue
        current.paragraphs.append(stripped)
    return sections


def parse_markdown(text: str) -> list[Section]:
    return parse_text(text)


def parse_docx(path: str) -> list[Section]:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise ParseFailed("缺少 python-docx 依赖") from exc
    try:
        d = docx.Document(path)
        sections: list[Section] = []
        current: Section | None = None
        heading_stack: list[str] = []
        for para in d.paragraphs:
            style_name = (para.style.name or "").lower() if para.style else ""
            text = para.text.strip()
            if not text:
                continue
            is_heading = "heading" in style_name or "标题" in style_name
            if is_heading:
                lvl_match = re.search(r"(\d)", style_name)
                level = int(lvl_match.group(1)) if lvl_match else 1
                heading_stack = heading_stack[: level - 1] + [text]
                current = _append(sections, list(heading_stack))
                continue
            if current is None:
                current = _append(sections, [])
            current.paragraphs.append(text)
        return sections
    except ParseFailed:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ParseFailed(f"DOCX 解析失败: {exc}") from exc


def parse_pdf(path: str) -> list[Section]:
    try:
        import fitz  # pymupdf
    except ImportError as exc:  # pragma: no cover
        raise ParseFailed("缺少 pymupdf 依赖") from exc
    try:
        doc = fitz.open(path)
        text = "\n".join(page.get_text("text") for page in doc)
        return parse_text(text)
    except Exception as exc:  # noqa: BLE001
        raise ParseFailed(f"PDF 解析失败: {exc}") from exc


def parse_file(path: str, file_name: str | None = None) -> list[Section]:
    """按扩展名路由解析。"""
    name = (file_name or path).lower()
    try:
        if name.endswith(".pdf"):
            return parse_pdf(path)
        if name.endswith(".docx"):
            return parse_docx(path)
        if name.endswith((".md", ".markdown")):
            with open(path, encoding="utf-8") as f:
                return parse_markdown(f.read())
        if name.endswith(".txt"):
            with open(path, encoding="utf-8") as f:
                return parse_text(f.read())
    except ParseFailed:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ParseFailed(f"解析失败（{name}）: {exc}") from exc
    raise ParseFailed(f"不支持的文件类型：{name}")
